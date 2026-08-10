from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


OUTPUT = Path(r"D:\work\workTicket3\output\国网安徽变电站第二种工作票_空白可编辑模板.docx")

PAGE_WIDTH_IN = 8.2677
PAGE_HEIGHT_IN = 11.6929
MARGIN_IN = 0.62
CONTENT_WIDTH_DXA = 10118
FONT_BODY = "宋体"
FONT_TITLE = "黑体"


def set_run_font(run, *, name=FONT_BODY, size=10, bold=None, color="000000"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    for key in ("ascii", "hAnsi", "eastAsia"):
        rpr.rFonts.set(qn(f"w:{key}"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def set_paragraph_format(paragraph, *, align=None, before=0, after=0, line=1.0, keep=False):
    if align is not None:
        paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep


def set_cell_margins(cell, top=35, start=55, bottom=35, end=55):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, spec in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key, value in spec.items():
            edge.set(qn(f"w:{key}"), str(value))


def set_table_borders(table, *, val="single", size=6, color="000000"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), val)
        if val != "nil":
            edge.set(qn("w:sz"), str(size))
            edge.set(qn("w:space"), "0")
            edge.set(qn("w:color"), color)


def set_table_geometry(table, widths):
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must total {CONTENT_WIDTH_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Twips(width)
            set_cell_margins(cell)


def set_row_min_height(row, points):
    row.height = Pt(points)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_text(cell, text="", *, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=10):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    set_paragraph_format(paragraph, align=align)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return paragraph


def set_section_geometry(section):
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)
    section.left_margin = Inches(MARGIN_IN)
    section.right_margin = Inches(MARGIN_IN)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)


def set_header_footer(section, page_number, *, show_header):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    hp = section.header.paragraphs[0]
    hp.clear()
    set_paragraph_format(hp, align=WD_ALIGN_PARAGRAPH.RIGHT)
    if show_header:
        run = hp.add_run("工作票编号：____________________________")
        set_run_font(run, size=8.5)

    fp = section.footer.paragraphs[0]
    fp.clear()
    set_paragraph_format(fp, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = fp.add_run(f"第{page_number}页    共4页")
    set_run_font(run, size=9)


def add_heading(doc, text, *, before=2, after=1, size=10.2):
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, before=before, after=after, keep=True)
    run = paragraph.add_run(text)
    set_run_font(run, size=size)
    return paragraph


def add_note(doc, text, *, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=0, line=1.0):
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, align=align, before=before, after=after, line=line)
    run = paragraph.add_run(text)
    set_run_font(run, size=size)
    return paragraph


def add_metadata_row(doc, labels, widths, *, height=17, size=9.5):
    table = doc.add_table(rows=1, cols=len(widths))
    set_table_geometry(table, widths)
    set_table_borders(table, val="nil")
    row = table.rows[0]
    set_row_min_height(row, height)
    for index, (label, is_value) in enumerate(labels):
        cell = row.cells[index]
        set_cell_text(cell, label, size=size)
        if is_value:
            set_cell_border(cell, bottom={"val": "single", "sz": "6", "space": "0", "color": "000000"})
    return table


def add_blank_lines(doc, count, *, min_height=17, size=9.5):
    table = doc.add_table(rows=count, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, val="nil")
    for row in table.rows:
        set_row_min_height(row, min_height)
        set_cell_text(row.cells[0], "", size=size)
        set_cell_border(row.cells[0], bottom={"val": "single", "sz": "4", "space": "0", "color": "000000"})
    return table


def add_two_column_form(doc, rows, widths, *, min_height=20, header=True, body_size=9.3):
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, widths)
    set_table_borders(table)
    for ri, values in enumerate(rows):
        set_row_min_height(table.rows[ri], min_height)
        for ci, value in enumerate(values):
            set_cell_text(
                table.cell(ri, ci),
                value,
                align=WD_ALIGN_PARAGRAPH.CENTER if header and ri == 0 else WD_ALIGN_PARAGRAPH.LEFT,
                bold=bool(header and ri == 0),
                size=9.2 if header and ri == 0 else body_size,
            )
    return table


def start_new_page(doc, page_number):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_geometry(section)
    set_header_footer(section, page_number, show_header=True)
    return section


def build_page_one(doc):
    title = doc.add_paragraph()
    set_paragraph_format(title, align=WD_ALIGN_PARAGRAPH.CENTER, before=1, after=1)
    run = title.add_run("国网安徽省电力有限公司变电站第二种工作票")
    set_run_font(run, name=FONT_TITLE, size=18.5, bold=True)

    level = doc.add_paragraph()
    set_paragraph_format(level, align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    run = level.add_run("________ 级")
    set_run_font(run, size=9.5)

    add_metadata_row(
        doc,
        [("单位：", False), ("", True), ("编号：", False), ("", True)],
        [760, 2500, 760, 6098],
    )
    add_metadata_row(
        doc,
        [("1.工作负责人（监护人）：", False), ("", True), ("班组：", False), ("", True)],
        [2700, 1200, 720, 5498],
    )

    add_heading(doc, "2.工作班人员（不包括工作负责人）", before=1, after=0)
    add_blank_lines(doc, 3, min_height=16)
    add_note(doc, "共 ______ 人", size=9.5, align=WD_ALIGN_PARAGRAPH.RIGHT)

    add_heading(doc, "3.工作的变、配电站名称及设备双重名称：", before=1, after=0)
    add_blank_lines(doc, 3, min_height=16)

    add_heading(doc, "4.工作任务：", before=1, after=1)
    task_rows = [("工作地点及设备双重名称", "工作内容")] + [("", "") for _ in range(11)]
    add_two_column_form(doc, task_rows, [5059, 5059], min_height=17.5, body_size=9)

    add_note(
        doc,
        "5.计划工作时间  自______年____月____日____时____分  至______年____月____日____时____分",
        size=9.2,
        before=1,
    )
    add_heading(doc, "6.工作条件（停电或不停电，或邻近及保留带电设备名称）：", before=1, after=0, size=9.6)
    add_blank_lines(doc, 1, min_height=16)
    add_heading(doc, "7.注意事项（安全措施）：", before=1, after=0)
    add_blank_lines(doc, 2, min_height=16)


def build_page_two(doc):
    add_note(doc, "7.注意事项（安全措施）（续）：", size=9.5, before=0, after=0)
    add_blank_lines(doc, 14, min_height=20)

    add_metadata_row(
        doc,
        [("工作票签发人签名：", False), ("", True), ("签发时间：", False), ("______年____月____日____时____分", True)],
        [2000, 2000, 1150, 4968],
        height=18,
    )
    add_metadata_row(
        doc,
        [("工作会签人签名：", False), ("", True), ("会签时间：", False), ("______年____月____日____时____分", True)],
        [2000, 2000, 1150, 4968],
        height=18,
    )

    add_heading(doc, "8.补充安全措施（工作许可人填写）：", before=2, after=0)
    add_blank_lines(doc, 2, min_height=18)

    add_heading(doc, "9.确认本工作票1—8项", before=2, after=0)
    add_note(doc, "许可工作时间：______年____月____日____时____分", size=9.5)
    add_metadata_row(doc, [("工作负责人签名：", False), ("", True), ("工作许可人签名：", False), ("", True)], [1650, 3200, 1650, 3618], height=18)

    add_heading(
        doc,
        "10.现场交底，工作班成员确认工作负责人布置的工作任务、人员分工、安全措施和注意事项并签名",
        before=2,
        after=0,
        size=9.5,
    )
    add_blank_lines(doc, 3, min_height=17)


def build_page_three(doc):
    add_note(doc, "10.现场交底签名（续）：", size=9.5, before=0, after=0)
    add_blank_lines(doc, 2, min_height=17)

    add_heading(doc, "11.工作票延期", before=2, after=0)
    add_note(doc, "有效期延长到 ______年____月____日____时____分", size=9.5)
    add_metadata_row(doc, [("工作负责人签名：", False), ("", True), ("签名时间：", False), ("______年____月____日____时____分", True)], [1700, 2300, 1150, 4968], height=17)
    add_metadata_row(doc, [("工作许可人签名：", False), ("", True), ("签名时间：", False), ("______年____月____日____时____分", True)], [1700, 2300, 1150, 4968], height=17)

    add_heading(doc, "12.工作负责人变动情况：", before=2, after=0)
    add_note(doc, "原工作负责人：____________ 离去，变更 ____________ 为工作负责人", size=9.5)
    add_note(doc, "工作票签发人：____________    ______年____月____日____时____分", size=9.5)

    add_heading(doc, "13.工作人员变动（变动人员姓名、变动日期及时间）", before=2, after=0)
    add_blank_lines(doc, 4, min_height=18)
    add_metadata_row(doc, [("工作负责人签名：", False), ("", True)], [1700, 8418], height=17)

    add_heading(doc, "14.每日开工和收工时间（使用一天的工作票不必填写）", before=2, after=1, size=9.8)
    build_daily_table(doc)

    add_heading(doc, "15.工作票终结", before=2, after=0)
    add_note(
        doc,
        "全部工作于 ______年____月____日____时____分 结束，工作人员已全部撤离，材料、工具已清理完毕。",
        size=9.3,
    )
    add_metadata_row(doc, [("工作负责人签名：", False), ("", True), ("工作许可人签名：", False), ("", True)], [1650, 3200, 1650, 3618], height=17)


def build_daily_table(doc):
    widths = [620, 620, 620, 620, 1289, 1289, 620, 620, 620, 620, 1289, 1291]
    table = doc.add_table(rows=8, cols=12)
    set_table_geometry(table, widths)
    set_table_borders(table)
    for row in table.rows:
        set_row_min_height(row, 20)

    stop = table.cell(0, 0).merge(table.cell(0, 3))
    set_cell_text(stop, "收工时间", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.3)
    start = table.cell(0, 6).merge(table.cell(0, 9))
    set_cell_text(start, "开工时间", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.3)
    for col, label in ((4, "工作负责人"), (5, "工作许可人"), (10, "工作许可人"), (11, "工作负责人")):
        merged = table.cell(0, col).merge(table.cell(1, col))
        set_cell_text(merged, label, align=WD_ALIGN_PARAGRAPH.CENTER, size=8.6)
    for col, label in enumerate(("月", "日", "时", "分")):
        set_cell_text(table.cell(1, col), label, align=WD_ALIGN_PARAGRAPH.CENTER, size=8.8)
        set_cell_text(table.cell(1, col + 6), label, align=WD_ALIGN_PARAGRAPH.CENTER, size=8.8)
    for row in table.rows[2:]:
        for cell in row.cells:
            set_cell_text(cell, "", align=WD_ALIGN_PARAGRAPH.CENTER, size=8.8)
    return table


def build_page_four(doc):
    add_heading(doc, "16.备注：", before=0, after=2, size=10.5)
    add_note(
        doc,
        "指定专责监护人：________________，负责监护：____________________________________________",
        size=9.7,
    )
    add_note(doc, "备注人：________________    ______年____月____日____时____分", size=9.7)
    add_blank_lines(doc, 12, min_height=22)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    rfonts = normal._element.rPr.rFonts
    for key in ("ascii", "hAnsi", "eastAsia"):
        rfonts.set(qn(f"w:{key}"), FONT_BODY)
    normal.font.size = Pt(10)
    fmt = normal.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0


def main():
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "国网安徽省电力有限公司变电站第二种工作票空白模板"
    doc.core_properties.subject = "依据真实工作票固定版式提取的可编辑空白模板"
    doc.core_properties.author = "项目实验数据构建"
    doc.core_properties.comments = "实例数据已清空；不得作为未经审核的现场工作票直接使用。"

    section = doc.sections[0]
    set_section_geometry(section)
    set_header_footer(section, 1, show_header=False)
    build_page_one(doc)

    start_new_page(doc, 2)
    build_page_two(doc)

    start_new_page(doc, 3)
    build_page_three(doc)

    start_new_page(doc, 4)
    build_page_four(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
