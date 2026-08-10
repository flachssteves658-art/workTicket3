from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def run_font_name(run):
    rpr = run._element.rPr
    if rpr is None or rpr.rFonts is None:
        return run.font.name
    return (
        rpr.rFonts.get(qn("w:eastAsia"))
        or rpr.rFonts.get(qn("w:ascii"))
        or run.font.name
    )


def paragraph_info(paragraph):
    return {
        "text": paragraph.text,
        "style": paragraph.style.name if paragraph.style else None,
        "alignment": int(paragraph.alignment) if paragraph.alignment is not None else None,
        "runs": [
            {
                "text": run.text,
                "font": run_font_name(run),
                "size_pt": run.font.size.pt if run.font.size else None,
                "bold": run.bold,
                "underline": run.underline,
            }
            for run in paragraph.runs
        ],
    }


def main():
    path = Path(sys.argv[1])
    doc = Document(path)
    result = {
        "paragraphs": [paragraph_info(p) for p in doc.paragraphs],
        "tables": [],
        "sections": [],
    }

    for s in doc.sections:
        result["sections"].append(
            {
                "page_width": s.page_width,
                "page_height": s.page_height,
                "left_margin": s.left_margin,
                "right_margin": s.right_margin,
                "top_margin": s.top_margin,
                "bottom_margin": s.bottom_margin,
                "header_distance": s.header_distance,
                "footer_distance": s.footer_distance,
            }
        )

    for ti, table in enumerate(doc.tables):
        table_data = {
            "index": ti,
            "rows": len(table.rows),
            "cols": len(table.columns),
            "grid_widths": [col.width for col in table.columns],
            "cells": [],
        }
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                table_data["cells"].append(
                    {
                        "row": ri,
                        "col": ci,
                        "text": cell.text,
                        "width": cell.width,
                        "paragraphs": [paragraph_info(p) for p in cell.paragraphs],
                    }
                )
        result["tables"].append(table_data)

    if len(sys.argv) > 2 and sys.argv[2] == "--compact":
        for i, paragraph in enumerate(doc.paragraphs):
            print(f"P{i}: {paragraph.text}")
        for ti, table in enumerate(doc.tables):
            print(f"TABLE {ti}: {len(table.rows)}x{len(table.columns)}")
            for ri, row in enumerate(table.rows):
                values = []
                seen = set()
                for ci, cell in enumerate(row.cells):
                    key = id(cell._tc)
                    if key in seen:
                        continue
                    seen.add(key)
                    values.append(f"C{ci}={cell.text!r}")
                print(f"  R{ri}: " + " | ".join(values))
        return
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
