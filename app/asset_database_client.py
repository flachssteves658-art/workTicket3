from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any
from urllib.parse import quote

import httpx

from app.config import settings
from app.ticket_integration import normalize_ticket_fields


EQUIPMENT_FIELDS = [
    "所属地市", "运维单位", "所属电站", "设备名称", "电压等级", "设备类型",
    "资产编号", "运行编号", "所属间隔", "设备(资产)状态", "运行状态", "调度命名",
    "安装位置", "设备编码", "功能位置编码",
]


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm(value: Any) -> str:
    text = _clean(value).upper().replace("Ⅲ", "III").replace("Ⅱ", "II").replace("Ⅰ", "I")
    return re.sub(r"[\s:：,，。;；()（）\-_/]+", "", text)


def _source_key(item: dict[str, str]) -> str:
    identity = "|".join(
        _clean(item.get(key))
        for key in ("所属电站", "设备编码", "资产编号", "设备名称", "运行编号")
    )
    return hashlib.sha256(identity.encode("utf-8")).hexdigest()


def _camera_ip(item: dict[str, str]) -> str:
    for key in ("IP地址", "摄像头IP", "ip_address", "IP"):
        value = _clean(item.get(key))
        if value:
            return value
    return ""


def _headers() -> dict[str, str]:
    if not settings.hazard_integration_token:
        raise RuntimeError("未配置 HAZARD_INTEGRATION_TOKEN")
    return {"X-Integration-Token": settings.hazard_integration_token}


def _url(path: str) -> str:
    return f"{settings.hazard_dbm_base_url}{path}"


def _read_catalog_files(camera_file: Path, equipment_file: Path) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，请执行 pip install -r requirements.txt") from exc

    camera_book = load_workbook(camera_file, read_only=True, data_only=True)
    camera_sheet = camera_book[camera_book.sheetnames[0]]
    camera_rows = list(camera_sheet.iter_rows(values_only=True))
    if len(camera_rows) < 3:
        raise ValueError("摄像头清单没有可读取的数据")
    camera_headers = [_clean(value) for value in camera_rows[1]]
    cameras: list[dict[str, str]] = []
    for row in camera_rows[2:]:
        item = {camera_headers[index]: _clean(value) for index, value in enumerate(row) if index < len(camera_headers)}
        if item.get("摄像头编号"):
            cameras.append(item)
    camera_book.close()

    equipment: list[dict[str, str]] = []
    with equipment_file.open("rb") as equipment_stream:
        is_ooxml = equipment_stream.read(2) == b"PK"
    if is_ooxml:
        with equipment_file.open("rb") as stream:
            equipment_book = load_workbook(stream, read_only=True, data_only=True)
            equipment_sheet = equipment_book[equipment_book.sheetnames[0]]
            rows = equipment_sheet.iter_rows(values_only=True)
            headers = [_clean(value) for value in next(rows)]
            indexes = {name: headers.index(name) for name in EQUIPMENT_FIELDS if name in headers}
            for row in rows:
                item = {name: _clean(row[column]) for name, column in indexes.items()}
                if item.get("设备名称"):
                    equipment.append(item)
            equipment_book.close()
    else:
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("该文件是真正的旧版XLS，缺少 xlrd，请执行 pip install -r requirements.txt") from exc
        book = xlrd.open_workbook(str(equipment_file), on_demand=True)
        equipment_sheet = book.sheet_by_index(0)
        headers = [_clean(equipment_sheet.cell_value(0, column)) for column in range(equipment_sheet.ncols)]
        indexes = {name: headers.index(name) for name in EQUIPMENT_FIELDS if name in headers}
        for row_index in range(1, equipment_sheet.nrows):
            item = {name: _clean(equipment_sheet.cell_value(row_index, column)) for name, column in indexes.items()}
            if item.get("设备名称"):
                equipment.append(item)
        book.release_resources()
    return cameras, equipment


async def import_catalogs(camera_file: Path, equipment_file: Path) -> dict[str, Any]:
    cameras, equipment = _read_catalog_files(camera_file, equipment_file)
    equipment_payload = []
    for item in equipment:
        search = " ".join(_clean(item.get(key)) for key in EQUIPMENT_FIELDS)
        equipment_payload.append(
            {
                "source_key": _source_key(item),
                "station_name": _clean(item.get("所属电站")),
                "equipment_name": _clean(item.get("设备名称")),
                "voltage_level": _clean(item.get("电压等级")),
                "equipment_type": _clean(item.get("设备类型")),
                "asset_no": _clean(item.get("资产编号")),
                "running_no": _clean(item.get("运行编号")),
                "interval_name": _clean(item.get("所属间隔")),
                "asset_status": _clean(item.get("设备(资产)状态")),
                "running_status": _clean(item.get("运行状态")),
                "dispatch_name": _clean(item.get("调度命名")),
                "install_location": _clean(item.get("安装位置")),
                "equipment_code": _clean(item.get("设备编码")),
                "functional_location_code": _clean(item.get("功能位置编码")),
                "normalized_search": _norm(search),
            }
        )
    camera_payload = []
    for item in cameras:
        channel = _clean(item.get("通道编号"))
        nvr = next(iter(re.findall(r"NVR\d+", channel.upper())), "")
        camera_payload.append(
            {
                "camera_id": _clean(item.get("摄像头编号")),
                "station_name": _clean(item.get("站点所属名称")),
                "nvr_name": nvr,
                "channel_no": channel,
                "install_location": _clean(item.get("安装位置")),
                "monitor_area": _clean(item.get("监控区域")),
                "ip_address": _camera_ip(item),
                "running_status": _clean(item.get("运行状态")),
                "retention_period": _clean(item.get("视频留存周期")),
            }
        )
    async with httpx.AsyncClient(timeout=120, trust_env=False) as client:
        response = await client.post(
            _url("/integration/catalog/import"),
            headers=_headers(),
            json={
                "equipment": equipment_payload,
                "cameras": camera_payload,
                "equipment_source": equipment_file.name,
                "camera_source": camera_file.name,
            },
        )
        response.raise_for_status()
        return response.json()


async def catalog_status() -> dict[str, Any]:
    try:
        async with httpx.AsyncClient(timeout=10, trust_env=False) as client:
            response = await client.get(_url("/integration/catalog/status"), headers=_headers())
            response.raise_for_status()
            return response.json()
    except httpx.HTTPStatusError as exc:
        message = (
            "两个项目的集成令牌不一致"
            if exc.response.status_code == 401
            else f"数据库管理接口返回 HTTP {exc.response.status_code}"
        )
        return {"ready": False, "storage": "mysql", "message": message}
    except Exception as exc:
        return {"ready": False, "storage": "mysql", "message": str(exc)}


async def match_ticket_assets(record: dict[str, Any]) -> dict[str, Any]:
    ticket_no = str(record.get("ticket_no") or record.get("ticket_id"))
    structured_data = normalize_ticket_fields(record.get("structured_data"))
    async with httpx.AsyncClient(timeout=120, trust_env=False) as client:
        response = await client.post(
            _url("/integration/tickets/match"),
            headers=_headers(),
            json={
                "ticket_no": ticket_no,
                "source_file": str(record.get("source_file") or ""),
                "structured_data": structured_data,
            },
        )
        response.raise_for_status()
        return response.json()


async def load_match_result(ticket_no: str) -> dict[str, Any] | None:
    async with httpx.AsyncClient(timeout=30, trust_env=False) as client:
        response = await client.get(
            _url(f"/integration/tickets/{quote(ticket_no, safe='')}/matches"),
            headers=_headers(),
        )
        if response.status_code == 404:
            return None
        response.raise_for_status()
        return response.json()


async def confirm_cameras(
    ticket_no: str,
    camera_ids: list[str],
    confirmed_by: str = "manual",
) -> dict[str, Any]:
    async with httpx.AsyncClient(timeout=30, trust_env=False) as client:
        response = await client.post(
            _url(
                f"/integration/tickets/{quote(ticket_no, safe='')}/confirm-cameras",
            ),
            headers=_headers(),
            json={"camera_ids": camera_ids, "confirmed_by": confirmed_by},
        )
        response.raise_for_status()
        return response.json()


async def list_stream_options() -> dict[str, Any]:
    async with httpx.AsyncClient(timeout=30, trust_env=False) as client:
        response = await client.get(
            _url("/integration/streams/options"),
            headers=_headers(),
        )
        response.raise_for_status()
        return response.json()


async def bind_camera_streams(
    bindings: list[dict[str, Any]],
    allow_station_mismatch: bool = False,
) -> dict[str, Any]:
    async with httpx.AsyncClient(timeout=30, trust_env=False) as client:
        response = await client.post(
            _url("/integration/cameras/bind-streams"),
            headers=_headers(),
            json={
                "bindings": bindings,
                "allow_station_mismatch": allow_station_mismatch,
            },
        )
        response.raise_for_status()
        return response.json()


async def load_detection_task(ticket_no: str) -> dict[str, Any] | None:
    async with httpx.AsyncClient(timeout=30, trust_env=False) as client:
        response = await client.get(
            _url(f"/integration/tickets/{quote(ticket_no, safe='')}/task"),
            headers=_headers(),
        )
        if response.status_code == 404:
            return None
        response.raise_for_status()
        return response.json()


async def cancel_detection_task(ticket_no: str) -> dict[str, Any] | None:
    async with httpx.AsyncClient(timeout=30, trust_env=False) as client:
        response = await client.post(
            _url(
                f"/integration/tickets/{quote(ticket_no, safe='')}/cancel-task",
            ),
            headers=_headers(),
        )
        if response.status_code == 404:
            return None
        response.raise_for_status()
        return response.json()
